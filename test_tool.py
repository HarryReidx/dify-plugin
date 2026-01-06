"""
测试脚本 - 验证 Excel 生成逻辑
注意：此脚本仅用于验证 openpyxl 逻辑，不包含 Dify SDK 调用
"""
import json
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def test_excel_generation():
    # 读取示例 JSON
    with open('example_checklist.json', 'r', encoding='utf-8') as f:
        checklist_data = json.load(f)
    
    table_name = checklist_data.get("tableName", "检查表")
    columns = checklist_data.get("columns", [])
    sections = checklist_data.get("sections", [])
    
    print(f"表名: {table_name}")
    print(f"列数: {len(columns)}")
    print(f"分组数: {len(sections)}")
    
    # 创建 Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "检查表"
    
    # 定义字体（全文使用宋体，大标题使用黑体）
    title_font = Font(name="黑体", bold=True, size=16)
    header_font = Font(name="宋体", bold=True, size=10.5)
    section_font = Font(name="宋体", bold=True, size=10.5)
    data_font = Font(name="宋体", size=10.5)
    
    # 定义填充色
    header_fill = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
    section_fill = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid")
    
    # 定义对齐方式
    title_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    section_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # 数据行对齐
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_indent_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    # 设置列宽（按照新要求）
    ws.column_dimensions['A'].width = 6      # 序号
    ws.column_dimensions['B'].width = 15     # 检查内容
    ws.column_dimensions['C'].width = 50     # 检查方式
    ws.column_dimensions['D'].width = 26     # 检查参考依据
    ws.column_dimensions['E'].width = 26     # 检查情况
    ws.column_dimensions['F'].width = 26     # 存在问题和建议
    ws.column_dimensions['G'].width = 4.5    # 得分
    
    current_row = 1
    
    # 写入表头
    max_col = max(7, len(columns))
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_col)
    cell = ws.cell(row=current_row, column=1, value=table_name)
    cell.font = title_font
    cell.alignment = title_alignment
    cell.border = thin_border
    ws.row_dimensions[current_row].height = 24  # 第1行行高24
    current_row += 1
    
    # 写入字段名
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[current_row].height = 30  # 第2行行高30
    current_row += 1
    
    # 写入数据
    total_items = 0
    for section in sections:
        section_title = section.get("title", "")
        items = section.get("items", [])
        
        if section_title:
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=max_col)
            cell = ws.cell(row=current_row, column=1, value=section_title)
            cell.font = section_font
            cell.fill = section_fill
            cell.alignment = section_alignment
            cell.border = thin_border
            # 不设置固定高度，自适应
            current_row += 1
        
        for item in items:
            item_id = item.get("id", "")
            content = item.get("content", "")
            method = item.get("method", "")
            reference = item.get("reference", "")
            situation = item.get("situation", "")
            suggestion = item.get("suggestion", "")
            score = item.get("score", "")
            
            row_data = [item_id, content, method, reference, situation, suggestion, score]
            while len(row_data) < len(columns):
                row_data.append("")
            
            for col_idx, value in enumerate(row_data[:len(columns)], start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                cell.font = data_font
                cell.border = thin_border
                
                # A, E, F, G列居中，B, C, D列左对齐+缩进
                if col_idx in [1, 5, 6, 7]:
                    cell.alignment = center_alignment
                else:
                    cell.alignment = left_indent_alignment
            
            # 不设置固定高度，让Excel根据内容自适应
            
            current_row += 1
            total_items += 1
    
    # 添加外边框
    last_row = current_row - 1
    last_col = max_col
    
    for row in range(1, last_row + 1):
        for col in range(1, last_col + 1):
            cell = ws.cell(row=row, column=col)
            
            left_style = 'medium' if col == 1 else 'thin'
            right_style = 'medium' if col == last_col else 'thin'
            top_style = 'medium' if row == 1 else 'thin'
            bottom_style = 'medium' if row == last_row else 'thin'
            
            cell.border = Border(
                left=Side(style=left_style, color='000000'),
                right=Side(style=right_style, color='000000'),
                top=Side(style=top_style, color='000000'),
                bottom=Side(style=bottom_style, color='000000')
            )
    
    # 保存文件
    output_filename = f"{table_name}_测试.xlsx"
    wb.save(output_filename)
    
    print(f"\n✓ Excel 文件生成成功！")
    print(f"✓ 文件名: {output_filename}")
    print(f"✓ 总行数: {current_row - 1}")
    print(f"✓ 检查项数量: {total_items}")
    print(f"\n请打开文件查看格式是否正确。")


if __name__ == "__main__":
    try:
        test_excel_generation()
    except Exception as e:
        print(f"✗ 错误: {str(e)}")
        import traceback
        traceback.print_exc()
